# -*- coding: utf-8 -*-
"""
Lecture du référentiel Excel
============================
Le classeur ``referentiel/referentiel_irrigation.xlsx`` est la source unique des
hypothèses. Il est relu automatiquement dès que sa date de modification change :
il suffit d'enregistrer le fichier puis de rafraîchir la page.

Chaque valeur est identifiée par une clé ``Feuille|Objet|code`` (par exemple
``Equipement|Pivot|capex_unitaire``). Les valeurs modifiées dans l'application
(paramètres marqués « Variable = oui ») sont transmises sous forme de
« surcharges » indexées par cette même clé.
"""

from __future__ import annotations

import os
import threading
import unicodedata
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

FEUILLES_PARAMETRES = ["Cultures", "Equipement", "Acces_ressource", "Ressources"]


def slug(texte: str) -> str:
    """Identifiant ASCII sans espace, utilisable dans un attribut HTML."""
    t = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")


def _texte(v) -> str:
    return "" if v is None else str(v).strip()


def _nombre(v) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").replace("\u202f", "").replace(" ", ""))
    except ValueError:
        return None


@dataclass
class Reference:
    id: str
    courte: str
    complete: str
    url: str = ""

    @property
    def citation(self) -> str:
        """Citation dans le texte, format APA : (Auteur, date)."""
        return f"({self.courte})"


@dataclass
class Param:
    feuille: str
    objet: str
    code: str
    libelle: str
    valeur: Optional[float]
    unite: str
    fiabilite: str
    variable: bool
    vmin: Optional[float]
    vmax: Optional[float]
    ref_id: str
    precision: str

    @property
    def cle(self) -> str:
        return f"{self.feuille}|{self.objet}|{self.code}"

    @property
    def id_html(self) -> str:
        return slug(self.cle)

    @property
    def est_pourcentage(self) -> bool:
        return self.unite.startswith("%")


@dataclass
class Referentiel:
    chemin: str
    date_modification: float = 0.0
    params: Dict[str, Param] = field(default_factory=dict)
    references: Dict[str, Reference] = field(default_factory=dict)
    cartes: Dict[str, Dict[str, str]] = field(default_factory=dict)
    exemples: List[Dict[str, Any]] = field(default_factory=list)
    departements: Dict[str, Dict[str, str]] = field(default_factory=dict)
    ordre: Dict[str, List[str]] = field(default_factory=dict)
    anomalies: List[str] = field(default_factory=list)

    # ------------------------------------------------------------------ accès
    def param(self, feuille: str, objet: str, code: str) -> Optional[Param]:
        return self.params.get(f"{feuille}|{objet}|{code}")

    def objets(self, feuille: str, exclure_general: bool = True) -> List[str]:
        return [o for o in self.ordre.get(feuille, []) if not (exclure_general and o == "Général")]

    def params_objet(self, feuille: str, objet: str) -> Dict[str, Param]:
        return {p.code: p for p in self.params.values() if p.feuille == feuille and p.objet == objet}

    @property
    def cultures(self) -> List[str]:
        return self.objets("Cultures")

    @property
    def equipements(self) -> List[str]:
        return self.objets("Equipement")

    @property
    def sources(self) -> List[str]:
        return self.objets("Acces_ressource")

    @property
    def agences(self) -> List[str]:
        return [o for o in self.objets("Ressources") if o != "Énergie"]

    def reference(self, ref_id: str) -> Optional[Reference]:
        return self.references.get(ref_id)

    def infobulle(self, p: Param) -> str:
        """Texte d'info-bulle : citation APA et précision."""
        ref = self.reference(p.ref_id)
        texte = ref.citation if ref else f"(réf. {p.ref_id or 'non renseignée'})"
        if p.precision:
            texte += f" — {p.precision}"
        return texte

    def valeur(self, feuille: str, objet: str, code: str, surcharges: Optional[Dict[str, Any]] = None,
               defaut: Optional[float] = None) -> Optional[float]:
        """Valeur du référentiel, remplacée par la surcharge si le paramètre est variable."""
        p = self.param(feuille, objet, code)
        if p is None:
            return defaut
        v = p.valeur
        if surcharges and p.cle in surcharges:
            s = _nombre(surcharges[p.cle])
            if s is not None and p.variable:
                if p.vmin is not None:
                    s = max(p.vmin, s)
                if p.vmax is not None:
                    s = min(p.vmax, s)
                v = s
        return defaut if v is None else v

    def agence_du_departement(self, code_dep: str) -> Optional[str]:
        d = self.departements.get(str(code_dep).upper().zfill(2))
        return d["agence"] if d else None

    def en_json(self) -> Dict[str, Any]:
        """Version sérialisable pour l'interface."""
        return {
            "params": {k: {"valeur": p.valeur, "unite": p.unite, "fiabilite": p.fiabilite,
                           "variable": p.variable, "min": p.vmin, "max": p.vmax, "libelle": p.libelle,
                           "infobulle": self.infobulle(p)} for k, p in self.params.items()},
            "cultures": self.cultures,
            "equipements": self.equipements,
            "sources": self.sources,
            "agences": self.agences,
            "departements": self.departements,
            "exemples": self.exemples,
        }


# =============================================================================
# Chargement
# =============================================================================
def _lire(chemin: str) -> Referentiel:
    wb = load_workbook(chemin, data_only=True, read_only=True)
    ref = Referentiel(chemin=chemin, date_modification=os.path.getmtime(chemin))

    # --- Références
    if "References" in wb.sheetnames:
        for ligne in wb["References"].iter_rows(min_row=2, values_only=True):
            if not ligne or not _texte(ligne[0]):
                continue
            ref.references[_texte(ligne[0])] = Reference(
                id=_texte(ligne[0]), courte=_texte(ligne[1]), complete=_texte(ligne[2]),
                url=_texte(ligne[3]) if len(ligne) > 3 else "")

    # --- Paramètres (format long)
    for feuille in FEUILLES_PARAMETRES:
        if feuille not in wb.sheetnames:
            ref.anomalies.append(f"Feuille « {feuille} » absente du référentiel.")
            continue
        ordre: List[str] = []
        for i, ligne in enumerate(wb[feuille].iter_rows(min_row=2, values_only=True), start=2):
            if not ligne or not _texte(ligne[0]) or not _texte(ligne[1]):
                continue
            objet, code = _texte(ligne[0]), _texte(ligne[1])
            valeur = _nombre(ligne[3])
            if valeur is None:
                ref.anomalies.append(f"{feuille}, ligne {i} : valeur manquante ou non numérique ({objet} / {code}).")
            fiab = _texte(ligne[5]).upper()[:1] or "E"
            p = Param(feuille=feuille, objet=objet, code=code, libelle=_texte(ligne[2]) or code,
                      valeur=valeur, unite=_texte(ligne[4]), fiabilite=fiab,
                      variable=_texte(ligne[6]).lower() in ("oui", "o", "yes", "1", "vrai"),
                      vmin=_nombre(ligne[7]), vmax=_nombre(ligne[8]), ref_id=_texte(ligne[9]),
                      precision=_texte(ligne[10]) if len(ligne) > 10 else "")
            if p.ref_id and p.ref_id not in ref.references:
                ref.anomalies.append(f"{feuille}, ligne {i} : référence « {p.ref_id} » absente de la feuille References.")
            ref.params[p.cle] = p
            if objet not in ordre:
                ordre.append(objet)
        ref.ordre[feuille] = ordre

    # --- Cartes
    if "Cartes" in wb.sheetnames:
        for ligne in wb["Cartes"].iter_rows(min_row=2, values_only=True):
            if not ligne or not _texte(ligne[0]):
                continue
            ref.cartes[_texte(ligne[0])] = {
                "code": _texte(ligne[1]).upper(),
                "besoin": _texte(ligne[2]), "rdt_irrigue": _texte(ligne[3]), "rdt_pluvial": _texte(ligne[4]),
                "unite_rdt": (_texte(ligne[5]) or "auto").lower(), "unite_besoin": (_texte(ligne[6]) or "auto").lower(),
                "scenario": _texte(ligne[7]), "crs": (_texte(ligne[8]) or "auto"),
            }

    # --- Exemples
    if "Exemples" in wb.sheetnames:
        for ligne in wb["Exemples"].iter_rows(min_row=2, values_only=True):
            if not ligne or ligne[0] is None or not _texte(ligne[1]):
                continue
            nb_u = _nombre(ligne[13])
            ref.exemples.append({
                "id": int(_nombre(ligne[0]) or 0), "titre": _texte(ligne[1]), "commune": _texte(ligne[2]),
                "departement": _texte(ligne[3]).zfill(2) if _texte(ligne[3]) else "",
                "lat": _nombre(ligne[4]), "lon": _nombre(ligne[5]), "culture": _texte(ligne[6]),
                "surface": _nombre(ligne[7]), "source": _texte(ligne[8]),
                "retenue": (_texte(ligne[9]) or "Individuelle").lower(),
                "nb_adherents": int(_nombre(ligne[10]) or 1), "zre": bool(_nombre(ligne[11]) or 0),
                "equipement": _texte(ligne[12]), "nb_unites": int(nb_u) if nb_u else None,
                "enjeu": _texte(ligne[14]),
            })

    # --- Départements
    if "Departements" in wb.sheetnames:
        for ligne in wb["Departements"].iter_rows(min_row=2, values_only=True):
            if not ligne or not _texte(ligne[0]) or not _texte(ligne[2]):
                continue
            code = _texte(ligne[0]).upper().zfill(2)
            ref.departements[code] = {"nom": _texte(ligne[1]), "agence": _texte(ligne[2]),
                                      "precision": _texte(ligne[3]) if len(ligne) > 3 else ""}
    wb.close()
    return ref


class GestionnaireReferentiel:
    """Garde le référentiel en mémoire et le relit si le fichier a changé."""

    def __init__(self, chemin: str):
        self.chemin = chemin
        self._ref: Optional[Referentiel] = None
        self._verrou = threading.Lock()
        self.derniere_erreur: Optional[str] = None

    def get(self) -> Referentiel:
        with self._verrou:
            try:
                mtime = os.path.getmtime(self.chemin)
            except OSError as e:
                if self._ref is None:
                    raise FileNotFoundError(f"Référentiel introuvable : {self.chemin}") from e
                return self._ref
            if self._ref is None or mtime != self._ref.date_modification:
                try:
                    self._ref = _lire(self.chemin)
                    self.derniere_erreur = None
                except Exception as e:  # fichier ouvert en écriture par Excel, etc.
                    self.derniere_erreur = f"Lecture du référentiel impossible : {e}"
                    if self._ref is None:
                        raise
            return self._ref

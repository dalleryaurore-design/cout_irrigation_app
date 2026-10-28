# -*- coding: utf-8 -*-
"""
Besoin en eau d'irrigation par la méthode SGPE
==============================================
Bilan hydrique journalier du sol dans la zone racinaire, sur le modèle du calendrier
d'irrigation de CROPWAT (Smith, 1992 ; FAO-56, chapitre 8) :

    Dr(j) = Dr(j−1) − Peff(j) + ETc(j) − I(j)          (Dr : épuisement de la réserve, mm)
    ETc   = Kc × ETP
    RU accessible  TAW(j) = RU du sol × min(Zr(j), profondeur de sol) ÷ profondeur de sol
    RU facilement utilisable  RAW(j) = p × TAW(j),  p = p_FAO + 0,04 × (5 − ETc), borné à [0,1 ; 0,8]
    irrigation  I(j) = Dr(j) dès que Dr(j) > RAW(j) (retour à la capacité au champ) ; en fin de
                cycle, I(j) est limité à ce qu'il faut pour atteindre la récolte sans stress
    besoin net (mm) = Σ I(j) sur la saison

- Pluie efficace : méthode USDA S.C.S. de CROPWAT sur le cumul mensuel
  (Peff = P × (125 − 0,2 P) / 125 si P ≤ 250 mm, 125 + 0,1 P au-delà), répartie
  uniformément sur les jours du mois ; l'excédent au-delà de la capacité au champ est perdu
  par drainage.
- Profondeur racinaire : de Zr_ini au semis à Zr_max (FAO-56, tableau 22) au pic de Kc,
  proportionnellement à la montée du Kc ; le sol nouvellement colonisé est à l'état initial.
- État hydrique initial : épuisement au semis, en % de la réserve accessible (0 % = capacité
  au champ, valeur par défaut de CROPWAT). Cultures pérennes (Kc > 0 toute l'année) : une
  année de mise en route, la seconde est comptée.
- Kc mensuels : feuille « Cultures » du référentiel (codes kc_01 à kc_12) ; Zr_max et p par
  culture (zr_max, p_fao) ; réserve utile, profondeur de sol, Zr_ini et épuisement initial
  dans l'objet « Général » de la même feuille.
- ETP et P mensuelles : ``climate_data.xlsx`` (DRIAS, EURO-CORDEX 2020, ALADIN63_CNRM-CM5,
  RCP4.5), une ligne par point de grille, horizon et mois. La valeur retenue pour la
  commune est celle du point de grille le plus proche de son centre.

La lecture du classeur prend plusieurs secondes : les données sont mises en cache
dans un fichier ``.npz`` voisin. Le cache est valide tant que l'empreinte SHA-256
du classeur ne change pas (la date de modification ne suffit pas : un checkout git
la réécrit et forcerait une relecture à chaque déploiement).
"""

from __future__ import annotations

import hashlib
import os
import threading
from typing import Any, Dict, List, Optional, Tuple

import numpy as np

HORIZONS = ["H1", "H2", "H3"]
JOURS_MOIS = (31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31)
# Valeurs de repli si le référentiel ne renseigne pas les paramètres du bilan hydrique
ZR_MAX_DEFAUT, P_DEFAUT, ZR_INI_DEFAUT, RU_DEFAUT, PROFONDEUR_SOL_DEFAUT = 1.0, 0.5, 0.25, 100.0, 2.0
LIBELLES_HORIZON = {"H1": "horizon proche (2021-2050)", "H2": "horizon moyen (2041-2070)",
                    "H3": "horizon lointain (2071-2100)"}
# Distance maximale entre le centre de la commune et le point de grille (grille de 8 km).
DISTANCE_MAX_KM = 12.0

_verrou = threading.Lock()
_donnees: Dict[str, Any] = {}


def _lire_classeur(chemin: str) -> Dict[str, np.ndarray]:
    from openpyxl import load_workbook
    wb = load_workbook(chemin, read_only=True, data_only=True)
    points: Dict[str, int] = {}
    coords: list = []
    valeurs: Dict[Tuple[int, int, int], Tuple[float, float]] = {}
    for ws in wb.worksheets:
        for ligne in ws.iter_rows(values_only=True):
            if not ligne or ligne[0] is None:
                continue
            s = str(ligne[0])
            if s.startswith("#"):
                continue
            p = s.split(";")
            if len(p) < 8 or p[4] not in HORIZONS:
                continue
            try:
                lat, lon, mois, rr, etp = float(p[1]), float(p[2]), int(p[5]), float(p[6]), float(p[7])
            except ValueError:
                continue
            if p[0] not in points:
                points[p[0]] = len(coords)
                coords.append((lat, lon))
            valeurs[(points[p[0]], HORIZONS.index(p[4]), mois - 1)] = (rr, etp)
    wb.close()
    n = len(coords)
    pluie = np.full((n, len(HORIZONS), 12), np.nan)
    etp = np.full((n, len(HORIZONS), 12), np.nan)
    for (i, h, m), (rr, e) in valeurs.items():
        pluie[i, h, m] = rr
        etp[i, h, m] = e
    return {"coords": np.array(coords, dtype="float64").reshape(-1, 2), "pluie": pluie, "etp": etp,
            "ids": np.array(list(points.keys()))}


def _empreinte(chemin: str) -> str:
    """SHA-256 du classeur. Indépendant de la date de modification du fichier."""
    h = hashlib.sha256()
    with open(chemin, "rb") as f:
        for bloc in iter(lambda: f.read(1 << 20), b""):
            h.update(bloc)
    return h.hexdigest()


def _empreinte_stockee(z) -> str:
    """Empreinte lue dans le npz. ``.item()`` évite qu'un scalaire NumPy ne se compare pas à la chaîne."""
    if "empreinte" not in z.files:
        return ""
    brut = z["empreinte"]
    if getattr(brut, "shape", None) == ():
        brut = brut.item()
    return str(brut).strip()


def _charger(chemin: str) -> Optional[Dict[str, Any]]:
    """Données climatiques en mémoire (depuis le cache .npz si le contenu n'a pas changé)."""
    if not chemin or not os.path.exists(chemin):
        return None
    mtime = os.path.getmtime(chemin)
    with _verrou:
        if _donnees.get("mtime") == mtime and _donnees.get("chemin") == chemin:
            return _donnees
        empreinte = _empreinte(chemin)
        cache = os.path.splitext(chemin)[0] + ".cache.npz"
        d = None
        if os.path.exists(cache):
            try:
                with np.load(cache) as z:
                    if _empreinte_stockee(z) == empreinte:
                        d = {k: np.array(z[k]) for k in ("coords", "pluie", "etp", "ids")}
            except Exception as exc:
                print(f"climat: cache illisible ({exc})", flush=True)
                d = None
        if d is None:
            print("climat: lecture du classeur (cache absent ou périmé)", flush=True)
            d = _lire_classeur(chemin)
            try:
                np.savez_compressed(cache, empreinte=np.array(empreinte), **d)
            except OSError:
                pass
        _donnees.clear()
        _donnees.update(d)
        _donnees.update({"mtime": mtime, "chemin": chemin})
        return _donnees


def prechauffer(chemin: str) -> None:
    """Charge le cache avant d'accepter des requêtes.

    Un thread en arrière-plan relisait le classeur en tenant le verrou : la première
    simulation attendait, et le seul process ne répondait plus à rien.
    """
    try:
        d = _charger(chemin)
        n = 0 if not d else len(d.get("coords", ()))
        print(f"climat prêt ({n} points)", flush=True)
    except Exception as exc:
        print(f"climat: chargement impossible ({exc})", flush=True)


def kc_mensuels(ref, culture: str, surcharges: Optional[Dict[str, Any]] = None) -> Optional[list]:
    """Les 12 Kc de la culture, ou None si le référentiel n'en contient pas."""
    kc = [ref.valeur("Cultures", culture, f"kc_{m:02d}", surcharges) for m in range(1, 13)]
    if all(k is None for k in kc):
        return None
    return [k or 0.0 for k in kc]


def point_proche(d: Dict[str, Any], lonlat: Tuple[float, float]) -> Tuple[int, float]:
    """Indice du point de grille le plus proche et distance en km."""
    lat0, lon0 = lonlat[1], lonlat[0]
    c = d["coords"]
    dy = (c[:, 0] - lat0) * 111.2
    dx = (c[:, 1] - lon0) * 111.2 * np.cos(np.radians(lat0))
    dist = np.hypot(dx, dy)
    i = int(np.argmin(dist))
    return i, float(dist[i])


def pluie_efficace(p: float) -> float:
    """Pluie efficace mensuelle (mm), méthode USDA S.C.S. de CROPWAT."""
    p = max(float(p), 0.0)
    return p * (125.0 - 0.2 * p) / 125.0 if p <= 250.0 else 125.0 + 0.1 * p


def _ordre_saison(kc: List[float]) -> Tuple[List[int], bool]:
    """Mois de la saison culturale dans l'ordre chronologique (semis d'abord), et culture pérenne ?

    Le semis est le premier mois où Kc > 0 après un mois où Kc = 0 (une culture d'hiver commence
    ainsi en octobre et se termine l'année suivante)."""
    if all(k > 0 for k in kc):
        return list(range(12)), True
    debut = next(m for m in range(12) if kc[m] > 0 and kc[m - 1] <= 0)
    mois = []
    for j in range(12):
        m = (debut + j) % 12
        if kc[m] <= 0:
            break
        mois.append(m)
    return mois, False


def bilan_hydrique(kc: List[float], etp: List[float], pluie: List[float], ru: float, profondeur_sol: float,
                   zr_max: float, zr_ini: float, p_fao: float, epuisement_initial: float,
                   journal: bool = False, reduire_fin_cycle: bool = True) -> Dict[str, Any]:
    """Bilan hydrique journalier (calendrier d'irrigation de CROPWAT), irrigation déclenchée
    à l'épuisement de la réserve facilement utilisable et ramenant le sol à la capacité au champ.

    ru : réserve utile du sol (mm) sur ``profondeur_sol`` (m) ; epuisement_initial : fraction (0-1)
    de la réserve accessible déjà consommée au semis. Renvoie le besoin net (mm), son détail
    mensuel et quelques indicateurs (nombre d'apports, ETc, pluie efficace, ETa sans irrigation).
    journal : ajoute le détail jour par jour de l'année comptée (validation, outils/comparer_modeles.py).
    reduire_fin_cycle : False reproduit CROPWAT, qui remplit le sol à la capacité au champ jusqu'à la récolte.
    """
    mois, perenne = _ordre_saison(kc)
    jours: List[Dict[str, float]] = []
    profondeur_sol = max(profondeur_sol, 0.05)
    ru_par_m = max(ru, 0.0) / profondeur_sol
    zr_max = max(min(zr_max, profondeur_sol), 0.05)
    zr_ini = max(min(zr_ini, zr_max), 0.05)
    e0 = max(min(epuisement_initial, 1.0), 0.0)
    kc_debut = kc[mois[0]]
    i_pic = max(range(len(mois)), key=lambda j: kc[mois[j]])
    kc_pic = kc[mois[i_pic]]

    def zr(j):
        if perenne or j >= i_pic or kc_pic <= kc_debut:
            return zr_max if (perenne or j >= i_pic) else zr_ini
        return zr_ini + (zr_max - zr_ini) * max(0.0, (kc[mois[j]] - kc_debut) / (kc_pic - kc_debut))

    irrig_m, etc_m, peff_m, eta_sec_m = [0.0] * 12, [0.0] * 12, [0.0] * 12, [0.0] * 12
    # Plus fort déficit cumulé à venir jusqu'à la récolte, jour par jour : max(0, max_k Σ (ETc − Peff) des
    # jours suivants jusqu'au jour k). Une pluie d'automne ne compense pas un déficit de l'été.
    nets = [kc[m] * etp[m] / JOURS_MOIS[m] - pluie_efficace(pluie[m]) / JOURS_MOIS[m]
            for m in mois for _ in range(JOURS_MOIS[m])]
    reste = [0.0] * (len(nets) + 1)
    for k in range(len(nets) - 1, -1, -1):
        reste[k] = max(0.0, nets[k] + reste[k + 1])
    n_apports = 0
    annees = 2 if perenne else 1   # pérenne : une année de mise en route
    z_prec = zr(0)
    dr = dr_sec = e0 * ru_par_m * z_prec
    for annee in range(annees):
        compte = annee == annees - 1
        jour = 0
        for j, m in enumerate(mois):
            n = JOURS_MOIS[m]
            eto, peff = etp[m] / n, pluie_efficace(pluie[m]) / n
            z = max(zr(j), z_prec)
            if z > z_prec:  # sol nouvellement colonisé par les racines, à l'état initial
                dr += e0 * ru_par_m * (z - z_prec)
                dr_sec += e0 * ru_par_m * (z - z_prec)
                z_prec = z
            taw = ru_par_m * z
            for _ in range(n):
                jour += 1
                etc = kc[m] * eto
                p = max(min(p_fao + 0.04 * (5.0 - etc), 0.8), 0.1)
                raw = p * taw
                # Conduite irriguée : pas de stress, apport dès que la RFU est épuisée. L'apport ramène le
                # sol à la capacité au champ, sauf en fin de cycle où il se limite à ce qu'il faut pour
                # atteindre la récolte sans stress (l'eau laissée dans le sol à la récolte n'est pas comptée).
                dr = max(dr - peff, 0.0) + etc
                if dr > raw:
                    dose = dr if not (compte and reduire_fin_cycle) else min(dr, max(dr + reste[jour] - raw, 0.0))
                    if compte and dose > 0:
                        irrig_m[m] += dose
                        n_apports += 1
                    dr -= dose
                # Conduite en sec : ETa réduite par le coefficient de stress Ks (FAO-56, éq. 84)
                dr_sec = max(dr_sec - peff, 0.0)
                ks = 1.0 if dr_sec <= raw else max((taw - dr_sec) / max((1.0 - p) * taw, 1e-9), 0.0)
                dr_sec = min(dr_sec + ks * etc, taw)
                if compte:
                    etc_m[m] += etc
                    peff_m[m] += peff
                    eta_sec_m[m] += ks * etc
                    if journal:
                        jours.append({"mois": m + 1, "kc": kc[m], "eto": eto, "etc": etc, "peff": peff,
                                      "zr": z, "taw": taw, "p": p, "raw": raw})
    total = sum(irrig_m)
    res = {"valeur": total, "mensuel": irrig_m, "etc": etc_m, "pluie_efficace": peff_m, "eta_sec": eta_sec_m,
           "n_apports": n_apports, "dose_moyenne": total / n_apports if n_apports else 0.0,
           "mois_saison": mois, "perenne": perenne, "taw_max": ru_par_m * zr_max, "zr_max": zr_max,
           "zr_ini": zr_ini, "satisfaction_sec": (sum(eta_sec_m) / sum(etc_m)) if sum(etc_m) else 1.0}
    if journal:
        res["journal"] = jours
    return res


def parametres_bilan(ref, culture: str, surcharges: Optional[Dict[str, Any]] = None) -> Dict[str, float]:
    """Paramètres du bilan hydrique lus dans le référentiel (feuille Cultures), surcharges comprises."""
    V = lambda o, c, d: ref.valeur("Cultures", o, c, surcharges, d)  # noqa: E731
    return {"zr_max": V(culture, "zr_max", ZR_MAX_DEFAUT), "p_fao": V(culture, "p_fao", P_DEFAUT),
            "zr_ini": V("Général", "zr_ini", ZR_INI_DEFAUT),
            "epuisement_initial": V("Général", "depletion_initiale", 0.0),
            "profondeur_sol": V("Général", "profondeur_sol", PROFONDEUR_SOL_DEFAUT),
            "ru_defaut": V("Général", "ru_defaut", RU_DEFAUT)}


def rendement_relatif(ky: float, satisfaction: float) -> float:
    """Rendement relatif Ya/Ym = 1 − Ky × (1 − ETa/ETc) (FAO-33), borné à [0 ; 1]."""
    return max(min(1.0 - ky * (1.0 - satisfaction), 1.0), 0.0)


def bilan_climatique(lonlat: Optional[Tuple[float, float]], chemin: str) -> Optional[Dict[str, Any]]:
    """Bilan climatique P − ETP au point de grille le plus proche, par horizon : déficit de la saison
    d'irrigation (avril à septembre) et excédent hivernal (somme des mois d'octobre à mars où P > ETP,
    eau disponible pour la recharge et le remplissage des retenues). None sans donnée à proximité."""
    d = _charger(chemin) if lonlat is not None else None
    if d is None or not len(d["coords"]):
        return None
    i, dist = point_proche(d, lonlat)
    if dist > DISTANCE_MAX_KM:
        return None
    ete, hiver = range(3, 9), (9, 10, 11, 0, 1, 2)
    res = {"point": str(d["ids"][i]), "distance_km": dist, "horizons": {}}
    for h, nom in enumerate(HORIZONS):
        p, e = d["pluie"][i, h], d["etp"][i, h]
        if not (np.all(np.isfinite(p)) and np.all(np.isfinite(e))):
            return None
        res["horizons"][nom] = {"p_ete": float(sum(p[m] for m in ete)), "etp_ete": float(sum(e[m] for m in ete)),
                                "bilan_ete": float(sum(p[m] - e[m] for m in ete)),
                                "excedent_hiver": float(sum(max(p[m] - e[m], 0.0) for m in hiver)),
                                "p_annuel": float(p.sum())}
    return res


def besoin_sgpe(ref, culture: str, lonlat: Optional[Tuple[float, float]], chemin: str,
                horizon: str = "H1", surcharges: Optional[Dict[str, Any]] = None,
                ru: Optional[float] = None, profondeur: Optional[float] = None) -> Dict[str, Any]:
    """Besoin en eau d'irrigation (mm) par la méthode SGPE (bilan hydrique journalier).

    ru : réserve utile du sol (mm, sur la profondeur de sol) ; profondeur : profondeur de sol (m) ;
    à défaut, valeurs par défaut du référentiel. Le détail donne aussi le rendement relatif en sec
    (Ky, FAO-33) lorsque la culture a un Ky.
    Renvoie {"valeur": float|None, "note": str, "detail": {...}} ; valeur None si le calcul
    n'est pas possible (pas de Kc, pas de commune, pas de donnée climatique à proximité).
    """
    kc = kc_mensuels(ref, culture, surcharges)
    if kc is None:
        return {"valeur": None, "note": f"Méthode SGPE indisponible pour « {culture} » : coefficients culturaux (Kc) "
                                        "non renseignés dans le référentiel."}
    if not any(k > 0 for k in kc):
        return {"valeur": None, "note": f"Méthode SGPE indisponible pour « {culture} » : tous les Kc sont nuls."}
    if lonlat is None:
        return {"valeur": None, "note": "Sélectionnez une commune pour appliquer la méthode SGPE."}
    d = _charger(chemin)
    if d is None or not len(d["coords"]):
        return {"valeur": None, "note": "Données climatiques (climate_data.xlsx) introuvables."}
    i, dist = point_proche(d, lonlat)
    if dist > DISTANCE_MAX_KM:
        return {"valeur": None, "note": "Pas de donnée climatique (ETP, P) à proximité de la commune "
                                        f"(point le plus proche à {dist:.0f} km)."}
    h = HORIZONS.index(horizon) if horizon in HORIZONS else 0
    etp, pluie = d["etp"][i, h], d["pluie"][i, h]
    if not (np.all(np.isfinite(etp)) and np.all(np.isfinite(pluie))):
        return {"valeur": None, "note": "Données climatiques incomplètes au point de grille le plus proche."}
    prm = parametres_bilan(ref, culture, surcharges)
    ru_retenue = ru if ru is not None and ru > 0 else prm["ru_defaut"]
    prof_retenue = min(profondeur, 2.0) if profondeur is not None and profondeur > 0 else prm["profondeur_sol"]
    b = bilan_hydrique(kc, [float(x) for x in etp], [float(x) for x in pluie], ru_retenue, prof_retenue,
                       prm["zr_max"], prm["zr_ini"], prm["p_fao"], prm["epuisement_initial"])
    ky = ref.valeur("Cultures", culture, "ky", surcharges)
    lat, lon = d["coords"][i]
    return {"valeur": float(b["valeur"]), "note": "",
            "detail": {"point": str(d["ids"][i]), "lat": float(lat), "lon": float(lon), "distance_km": dist,
                       "horizon": HORIZONS[h], "libelle_horizon": LIBELLES_HORIZON[HORIZONS[h]],
                       "kc": kc, "etp": [float(x) for x in etp], "pluie": [float(x) for x in pluie],
                       "ru": ru_retenue, "profondeur_sol": prof_retenue, "parametres": prm, "deficit": b["mensuel"],
                       "ky": ky, "rendement_relatif_sec": (rendement_relatif(ky, b["satisfaction_sec"])
                                                           if ky is not None else None),
                       **{k: b[k] for k in ("etc", "pluie_efficace", "eta_sec", "n_apports", "dose_moyenne",
                                            "perenne", "taw_max", "zr_max", "satisfaction_sec")}}}
